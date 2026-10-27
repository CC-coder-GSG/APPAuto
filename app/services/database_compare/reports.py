from __future__ import annotations

import json
import re
import sqlite3
import uuid
import zipfile
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from .result import read_tables
from .values import dumps


def _json_report(result: Path, target: Path, summary: dict):
    with closing(sqlite3.connect(result.resolve().as_uri() + "?mode=ro", uri=True)) as db, target.open("w", encoding="utf-8") as out:
        out.write('{"summary":' + dumps(summary) + ',"tables":' + dumps(read_tables(result)) + ',"differences":[')
        first = True
        for row in db.execute("SELECT id,database_name,table_name,scope,kind,details_json FROM differences ORDER BY id"):
            if not first:
                out.write(",")
            first = False
            out.write(dumps({"id": row[0], "database": row[1], "table": row[2], "scope": row[3], "kind": row[4],
                             "details": json.loads(row[5])}))
        out.write("]}")


def _cells(sheet, values, reference=""):
    cells = []
    for value in values:
        if value is None:
            value = ""
        if not isinstance(value, (str, int, float, bool)):
            value = dumps(value)
        if isinstance(value, str):
            value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda m: f"\\u{ord(m[0]):04x}", value)
            if len(value) > 32000:
                value = value[:31000] + f"\n[显示节选，完整内容见 complete.json {reference}]"
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = "s"
            cells.append(cell)
        else:
            cells.append(value)
    return cells


def export_report(result: Path, directory: Path, summary: dict, format: str) -> Path:
    if format not in ("json", "excel"):
        raise ValueError("仅支持 json 或 excel")
    directory.mkdir(parents=True, exist_ok=True)
    identity = uuid.uuid4().hex
    complete = directory / f"{identity}.json"
    _json_report(result, complete, summary)
    if format == "json":
        return complete
    workbook = Workbook(write_only=True)
    overview = workbook.create_sheet("比较摘要")
    overview.append(["项目", "内容"])
    for key, value in summary.items():
        overview.append(_cells(overview, [key, value], reference=f"summary.{key}"))
    table_sheet = workbook.create_sheet("表覆盖与统计")
    table_sheet.append(["数据库", "表", "状态", "匹配方式", "旧记录", "新记录", "新增", "删除", "修改", "完整信息"])
    for table in read_tables(result):
        table_sheet.append(_cells(table_sheet, [table["database"], table["table"], table["status"], table.get("match_mode"),
            table.get("old_rows"), table.get("new_rows"), table.get("rows_added"), table.get("rows_removed"),
            table.get("rows_modified"), table], reference="tables"))
    headers = ["差异编号", "数据库", "表/对象", "范围", "差异类型", "字段/属性", "记录定位", "旧值", "新值", "完整明细"]
    sheet = None
    with closing(sqlite3.connect(result.resolve().as_uri() + "?mode=ro", uri=True)) as db:
        for index, row in enumerate(db.execute("SELECT id,database_name,table_name,scope,kind,details_json FROM differences ORDER BY id")):
            if index % 1000000 == 0:
                sheet = workbook.create_sheet(f"差异明细{index // 1000000 + 1}")
                sheet.append(headers)
            detail = json.loads(row[5])
            sheet.append(_cells(sheet, [*row[:5], detail.get("field", detail.get("property", "")),
                detail.get("key", ""), detail.get("old"), detail.get("new"), detail], reference=f"differences.id={row[0]}"))
    if sheet is None:
        workbook.create_sheet("差异明细1").append(headers)
    xlsx = directory / f"{identity}.xlsx"
    workbook.save(xlsx)
    archive = directory / f"{identity}.zip"
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as package:
        package.write(xlsx, "数据库对比报告.xlsx")
        package.write(complete, "complete.json")
    # These are exact files created by this invocation; the ZIP retains both.
    xlsx.unlink()
    complete.unlink()
    return archive
