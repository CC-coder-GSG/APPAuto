from __future__ import annotations

import json
import sqlite3
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.database_compare.engine import compare_databases
from app.services.database_compare.reports import export_report
from app.services.database_compare.result import read_result, read_tables
from app.services.database_compare.schema import fingerprint


def database(path, sql, rows=(), table="sample"):
    with sqlite3.connect(path) as db:
        db.executescript(sql)
        if rows:
            db.executemany(f'INSERT INTO "{table}" VALUES({",".join("?" for _ in rows[0])})', rows)
    return path


def compare(tmp_path, old, new, rules=None, **kwargs):
    result = tmp_path / "result.sqlite"
    summary = compare_databases([{"name": "sample.db", "old": old, "new": new}], result, rules, **kwargs)
    return summary, result, read_result(result, page_size=1000)["items"]


def test_order_independent_and_readonly(tmp_path):
    sql = "CREATE TABLE sample(id INTEGER, value TEXT)"
    a = database(tmp_path / "a.db", sql, [(1,"one"),(2,"two"),(2,"two")])
    b = database(tmp_path / "b.db", sql, [(2,"two"),(1,"one"),(2,"two")])
    before = [fingerprint(p) for p in (a,b)]
    summary, result, diffs = compare(tmp_path,a,b)
    assert summary["conclusion"] == "equal" and diffs == []
    assert [fingerprint(p) for p in (a,b)] == before
    assert read_tables(result)[0]["match_mode"] == "row_multiset"
    assert not list(tmp_path.glob("*.db-wal")) and not list(tmp_path.glob("*.db-shm"))


def test_keyed_changes_and_type_sensitive_values(tmp_path):
    sql = "CREATE TABLE sample(id INTEGER PRIMARY KEY, value BLOB, binary BLOB)"
    a = database(tmp_path/"a.db",sql,[(1,1,b"abc"),(2,None,b"abc"),(3,"",b"")])
    b = database(tmp_path/"b.db",sql,[(1,1.0,b"abd"),(2,"",b"abc"),(4,0,b"")])
    summary,_,diffs = compare(tmp_path,a,b)
    assert summary["status"] == "succeeded"
    assert (summary["rows_added"],summary["rows_removed"],summary["rows_modified"],summary["fields_modified"]) == (1,1,2,3)
    values = [d["details"] for d in diffs if d["kind"] == "value_changed"]
    assert any(v["old"]["type"] == "integer" and v["new"]["type"] == "real" for v in values)
    assert any(v["old"]["type"] == "blob" and v["old"]["length"] == v["new"]["length"] for v in values)
    assert any(v["old"]["type"] == "null" and v["new"]["type"] == "text" for v in values)


def test_duplicate_multiplicity_not_deduplicated(tmp_path):
    sql = "CREATE TABLE sample(ID INTEGER, value TEXT)"
    a = database(tmp_path/"a.db",sql,[(5,"a"),(5,"a"),(5,"a"),(5,"b")])
    b = database(tmp_path/"b.db",sql,[(5,"a"),(5,"b")])
    summary,_,diffs = compare(tmp_path,a,b)
    assert summary["rows_removed"] == 2
    assert len(diffs) == 1
    assert diffs[0]["details"]["old_count"] == 3 and diffs[0]["details"]["new_count"] == 1


def test_explicit_duplicate_and_null_keys_fail_without_false_equality(tmp_path):
    sql = "CREATE TABLE sample(id INTEGER,value TEXT)"
    a = database(tmp_path/"a.db",sql,[(1,"a"),(1,"b")])
    b = database(tmp_path/"b.db",sql,[(1,"a"),(None,"b")])
    rules = {"tables":{"sample.db.sample":{"match_mode":"key","key":["id"]}}}
    summary,result,diffs = compare(tmp_path,a,b,rules)
    assert summary["status"] == "partial" and summary["conclusion"] == "inconclusive"
    assert read_tables(result)[0]["status"] == "failed"
    assert [d["kind"] for d in diffs] == ["table_error"]


def test_nullable_declared_primary_key_falls_back_explicitly(tmp_path):
    sql = "CREATE TABLE sample(id TEXT PRIMARY KEY,value TEXT)"
    a = database(tmp_path/"a.db",sql,[(None,"a"),(None,"a")])
    b = database(tmp_path/"b.db",sql,[(None,"a")])
    summary,result,_ = compare(tmp_path,a,b)
    table = read_tables(result)[0]
    assert summary["rows_removed"] == 1 and table["match_mode"] == "row_multiset"
    assert any("NULL" in n for n in table["notes"])


def test_schema_constraints_indexes_views_triggers_and_generated_columns(tmp_path):
    a = database(tmp_path/"a.db", """
        CREATE TABLE sample(id INTEGER PRIMARY KEY, value TEXT DEFAULT 'a' CHECK(length(value)>0));
        CREATE VIEW v AS SELECT id FROM sample;
        CREATE INDEX idx ON sample(value);
    """)
    b = database(tmp_path/"b.db", """
        CREATE TABLE sample(id INTEGER PRIMARY KEY, value TEXT NOT NULL DEFAULT 'b' CHECK(length(value)>1),
                            derived TEXT GENERATED ALWAYS AS (value || '!') VIRTUAL);
        CREATE VIEW v AS SELECT value FROM sample;
        CREATE UNIQUE INDEX idx ON sample(value COLLATE NOCASE DESC) WHERE value != '';
        CREATE TRIGGER tr AFTER INSERT ON sample BEGIN SELECT 1; END;
    """)
    summary,_,diffs = compare(tmp_path,a,b)
    assert summary["status"] == "succeeded" and summary["has_changes"]
    assert any(d["kind"] == "column_added" and d["details"]["field"] == "derived" for d in diffs)
    assert {"default","not_null"} <= {d["details"].get("property") for d in diffs}
    assert {"view","trigger","index"} <= {d["details"].get("object_type") for d in diffs}


def test_simultaneous_schema_and_data_changes_with_key(tmp_path):
    a = database(tmp_path/"a.db","CREATE TABLE sample(id INTEGER PRIMARY KEY,value TEXT,removed TEXT)",[(1,"a","old")])
    b = database(tmp_path/"b.db","CREATE TABLE sample(id INTEGER PRIMARY KEY,value TEXT,added TEXT)",[(1,"b","new")])
    summary,_,diffs = compare(tmp_path,a,b)
    assert summary["status"] == "succeeded" and summary["rows_modified"] == 1
    assert {"value_changed","value_added","value_removed"} <= {d["kind"] for d in diffs}


def test_no_key_schema_change_retains_side_evidence_and_marks_partial(tmp_path):
    a = database(tmp_path/"a.db","CREATE TABLE sample(value TEXT,removed TEXT)",[("a","old"),("a","other")])
    b = database(tmp_path/"b.db","CREATE TABLE sample(value TEXT,added TEXT)",[("a","new")])
    summary,result,diffs = compare(tmp_path,a,b)
    assert summary["status"] == "partial" and summary["conclusion"] != "equal"
    assert len([d for d in diffs if d["scope"] == "evidence"]) == 3
    assert read_tables(result)[0]["new_rows"] == 1


def test_database_and_table_additions_preserve_rows(tmp_path):
    b = database(tmp_path/"b.db","CREATE TABLE sample(id INTEGER,value TEXT)",[(1,"x"),(1,"x")])
    summary,_,diffs = compare(tmp_path,None,b)
    assert summary["rows_added"] == 2
    assert {"database_added","table_added","row_added"} <= {d["kind"] for d in diffs}


@pytest.mark.parametrize("case",["corrupt","wal","virtual"])
def test_unsupported_or_corrupt_never_equal(tmp_path,case):
    a = database(tmp_path/"a.db","CREATE TABLE sample(id INTEGER)")
    b = database(tmp_path/"b.db","CREATE TABLE sample(id INTEGER)")
    if case == "corrupt":
        b.write_bytes(b"SQLite format 3\x00" + b"bad"*100)
    elif case == "wal":
        Path(str(b)+"-wal").write_bytes(b"uncheckpointed")
    else:
        with sqlite3.connect(b) as db:
            db.execute("CREATE VIRTUAL TABLE spatial USING rtree(id,minx,maxx,miny,maxy)")
    summary,_,_ = compare(tmp_path,a,b)
    assert summary["status"] in ("failed","partial") and summary["conclusion"] != "equal"
    assert summary["errors"]


def test_cancellation_retains_completed_coverage(tmp_path):
    a = database(tmp_path/"a.db","CREATE TABLE first(id INTEGER); CREATE TABLE second(id INTEGER)")
    flag = [False]
    def progress(data):
        if data["tables_completed"] == 1:
            flag[0] = True
    summary,_,_ = compare(tmp_path,a,a,progress=progress,cancelled=lambda:flag[0])
    assert summary["status"] == "cancelled" and summary["conclusion"] == "inconclusive"
    assert summary["tables_completed"] == 1 and summary["tables_total"] == 2


def test_source_change_during_comparison_invalidates_result(tmp_path):
    a = database(tmp_path/"a.db","CREATE TABLE sample(id INTEGER)")
    changed = [False]
    def progress(data):
        if data["tables_completed"] == 1 and not changed[0]:
            with sqlite3.connect(a) as db:
                db.execute("INSERT INTO sample VALUES(1)")
            changed[0] = True
    summary,_,_ = compare(tmp_path,a,a,progress=progress)
    assert summary["status"] == "partial" and summary["conclusion"] != "equal"


def test_report_preserves_long_values_and_formula_text(tmp_path):
    sql = "CREATE TABLE sample(id INTEGER PRIMARY KEY,value TEXT)"
    a = database(tmp_path/"a.db",sql,[(1,"old")])
    text = '=SUM(1,2)' + '\x01' + '文本'*18000
    b = database(tmp_path/"b.db",sql,[(1,text)])
    summary,result,_ = compare(tmp_path,a,b)
    preview = read_result(result,preview=True)
    assert preview["items"][0]["preview_only"] is True
    assert len(preview["items"][0]["details"]["new"]["value"]) < len(text)
    path = export_report(result,tmp_path/"exports",summary,"excel")
    with zipfile.ZipFile(path) as archive:
        complete = json.loads(archive.read("complete.json"))
        assert complete["differences"][0]["details"]["new"]["value"] == text
        xlsx = tmp_path/"report.xlsx"
        xlsx.write_bytes(archive.read("数据库对比报告.xlsx"))
    workbook = load_workbook(xlsx,read_only=True)
    try:
        rows = list(workbook["差异明细1"].iter_rows())
        assert "complete.json" in rows[1][-1].value
        assert rows[1][-1].data_type == "s"
    finally:
        workbook.close()


def test_unknown_rules_and_output_overwrite_rejected(tmp_path):
    a = database(tmp_path/"a.db","CREATE TABLE sample(id INTEGER)")
    with pytest.raises(ValueError):
        compare(tmp_path,a,a,{"ignored_tables":["sample"]})
    with pytest.raises(ValueError):
        compare_databases([{"name":"a","old":a,"new":a}],a)
    summary,_,_ = compare(tmp_path,a,a)
    assert summary["status"] == "succeeded"
    with pytest.raises(ValueError):
        compare(tmp_path,a,a)
